"""Parse GST portal GSTR-2A / 2B Excel or CSV exports into normalized rows."""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "supplier_gstin": ("gstin of supplier", "supplier gstin", "gstin", "ctin", "gstin_of_supplier"),
    "supplier_name": ("trade/legal name", "trade name", "legal name", "supplier name", "name of supplier"),
    "invoice_number": ("invoice number", "invoice no", "inum", "doc no", "document number", "invoice#"),
    "invoice_date": ("invoice date", "invoice dt", "idt", "document date", "doc date"),
    "taxable_amount": ("taxable value", "taxable amt", "taxable amount", "txval"),
    "cgst_amount": ("cgst", "central tax", "cgst amount", "camt"),
    "sgst_amount": ("sgst", "state tax", "sgst amount", "samt", "utgst"),
    "igst_amount": ("igst", "integrated tax", "igst amount", "iamt"),
    "net_amount": ("invoice value", "total invoice value", "total value", "net amount", "grand total"),
    "place_of_supply": ("place of supply", "pos"),
    "document_type": ("invoice type", "document type", "doc type", "inv_typ"),
}


def parse_gstr_file(filename: str, content: bytes) -> list[dict[str, Any]]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _parse_csv(content)
    if name.endswith((".xlsx", ".xls")):
        return _parse_xlsx(content)
    # Try CSV first (many portal "Excel" downloads are CSV), then xlsx.
    try:
        return _parse_csv(content)
    except Exception:
        return _parse_xlsx(content)


def _parse_csv(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    return _rows_from_table(rows)


def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel imports") from exc
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = wb.active
    rows = [[cell if cell is not None else "" for cell in row] for row in sheet.iter_rows(values_only=True)]
    return _rows_from_table(rows)


def _rows_from_table(rows: list[list[Any]]) -> list[dict[str, Any]]:
    header_idx, mapping = _find_header(rows)
    if header_idx is None:
        raise ValueError(
            "Could not find GSTR header row. Expected columns like "
            "'GSTIN of supplier', 'Invoice number', 'Invoice Date'."
        )
    parsed: list[dict[str, Any]] = []
    seen: set[tuple] = set()
    for raw in rows[header_idx + 1 :]:
        if not raw or all(str(c).strip() == "" for c in raw if c is not None):
            continue
        item = _map_row(raw, mapping)
        if not item.get("supplier_gstin") and not item.get("invoice_number"):
            continue
        key = (
            item.get("supplier_gstin"),
            (item.get("invoice_number") or "").upper(),
            item.get("invoice_date"),
        )
        if key in seen:
            continue
        seen.add(key)
        parsed.append(item)
    return parsed


def _find_header(rows: list[list[Any]]) -> tuple[int | None, dict[str, int]]:
    for idx, row in enumerate(rows[:40]):
        labels = [_norm_header(c) for c in row]
        mapping: dict[str, int] = {}
        for field, aliases in HEADER_ALIASES.items():
            for col, label in enumerate(labels):
                if label in aliases or any(a in label for a in aliases if len(a) > 4):
                    mapping[field] = col
                    break
        if "supplier_gstin" in mapping and "invoice_number" in mapping:
            return idx, mapping
    return None, {}


def _norm_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def _map_row(raw: list[Any], mapping: dict[str, int]) -> dict[str, Any]:
    def cell(field: str) -> Any:
        col = mapping.get(field)
        if col is None or col >= len(raw):
            return None
        return raw[col]

    taxable = _decimal(cell("taxable_amount"))
    cgst = _decimal(cell("cgst_amount"))
    sgst = _decimal(cell("sgst_amount"))
    igst = _decimal(cell("igst_amount"))
    net = _decimal(cell("net_amount"))
    if net is None and taxable is not None:
        net = taxable + (cgst or Decimal("0")) + (sgst or Decimal("0")) + (igst or Decimal("0"))

    gstin = _gstin(cell("supplier_gstin"))
    return {
        "supplier_gstin": gstin,
        "supplier_name": _str(cell("supplier_name")),
        "invoice_number": _str(cell("invoice_number")),
        "invoice_date": _date(cell("invoice_date")),
        "taxable_amount": taxable,
        "cgst_amount": cgst,
        "sgst_amount": sgst,
        "igst_amount": igst,
        "net_amount": net,
        "place_of_supply": _str(cell("place_of_supply")),
        "document_type": _str(cell("document_type")),
        "raw_json": {str(i): ("" if v is None else str(v)) for i, v in enumerate(raw)},
    }


def _str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _gstin(value: Any) -> str | None:
    text = _str(value)
    if not text:
        return None
    compact = re.sub(r"[^A-Za-z0-9]", "", text).upper()
    return compact or None


def _decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).replace(",", "").replace("₹", "").strip()
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None
